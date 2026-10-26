# -*- coding: utf-8 -*-
"""gongwen_docx 的回归测试。

不依赖真实字体：凡是要生成文档的用例，都先把字体检查打桩成「齐全」，
这样 Linux / macOS 的 CI 也能跑。真正跑字体检测的用例只测分支逻辑。
"""
import contextlib
import io
import os
import shutil
import struct
import sys
import tempfile
import types
import unittest
from unittest import mock

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.join(os.path.dirname(HERE), "scripts")
sys.path.insert(0, SCRIPTS)

import gongwen_docx as gw  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Mm  # noqa: E402


def _args(**kw):
    base = dict(out="a.docx", title="t", date=None, to=None, body=None, attach=None,
                sign=None, note=None, hide_first_page_number=False)
    base.update(kw)
    return types.SimpleNamespace(**base)


class ParseBodyTest(unittest.TestCase):
    def test_blank_line_splits(self):
        self.assertEqual(gw.parse_body("第一行\n第二行\n\n第二段\n"),
                         [("p", "第一行第二行"), ("p", "第二段")])

    def test_no_blank_line_falls_back_to_one_line_per_par(self):
        blocks = gw.parse_body("甲\n乙\n丙\n")
        self.assertEqual([b[1] for b in blocks], ["甲", "乙", "丙"])

    def test_heading_levels(self):
        blocks = gw.parse_body("# 一、\n## （一）\n### 1.\n#### （1）\n")
        self.assertEqual([b[0] for b in blocks], ["h1", "h2", "h3", "h4"])
        self.assertEqual([b[1] for b in blocks], ["一、", "（一）", "1.", "（1）"])

    def test_hash_escape(self):
        self.assertEqual(gw.parse_body("\\# 正文里的井号\n"), [("p", "# 正文里的井号")])

    def test_ascii_join_adds_one_space(self):
        self.assertEqual(gw.parse_body("用 DeepSeek\nHarness 测试\n\n尾段\n"),
                         [("p", "用 DeepSeek Harness 测试"), ("p", "尾段")])

    def test_fullwidth_digits_do_not_get_extra_space(self):
        # 全角数字、全角百分号都不是 ASCII，合并时不该插空格
        self.assertEqual(gw.parse_body("第１２３项\n合计５０％\n\n尾段\n"),
                         [("p", "第１２３项合计５０％"), ("p", "尾段")])

    def test_ascii_next_to_chinese_adds_no_space(self):
        # 只有两侧都是 ASCII 字母数字才补空格，「组」+「每」不补
        self.assertEqual(gw.parse_body("共 3 组\n每组 8 次\n\n尾段\n"),
                         [("p", "共 3 组每组 8 次"), ("p", "尾段")])

    def test_chinese_join_has_no_space(self):
        self.assertEqual(gw.parse_body("第一行\n第二行\n\n尾段\n"),
                         [("p", "第一行第二行"), ("p", "尾段")])

    def test_dash_list_items_each_become_one_paragraph(self):
        blocks = gw.parse_body("- 第一条\n- 第二条\n\n尾段\n")
        self.assertEqual(blocks, [("p", "第一条"), ("p", "第二条"), ("p", "尾段")])

    def test_numbered_list_items_each_become_one_paragraph(self):
        blocks = gw.parse_body("1. 甲\n2. 乙\n\n尾段\n")
        self.assertEqual([b[1] for b in blocks], ["甲", "乙", "尾段"])

    def test_paren_list_markers_are_stripped(self):
        blocks = gw.parse_body("（1）甲\n（一）乙\n\n尾段\n")
        self.assertEqual([b[1] for b in blocks], ["甲", "乙", "尾段"])

    def test_list_items_work_without_any_blank_line(self):
        blocks = gw.parse_body("- 甲\n- 乙\n")
        self.assertEqual([b[1] for b in blocks], ["甲", "乙"])

    def test_decimal_is_not_mistaken_for_a_list_item(self):
        # 「1.5」的「1.」后面跟的是数字，不是列表项前缀
        self.assertEqual(gw.parse_body("1.5 倍行距\n2.0 倍上限\n\n尾段\n"),
                         [("p", "1.5 倍行距2.0 倍上限"), ("p", "尾段")])

    def test_list_marker_escape_keeps_the_marker(self):
        self.assertEqual(gw.parse_body("\\- 这不是列表\n"), [("p", "- 这不是列表")])

    def test_trailing_backslash_forces_its_own_paragraph(self):
        self.assertEqual(gw.parse_body("Yours sincerely,\\\nLi Ming\n\n尾段\n"),
                         [("p", "Yours sincerely,"), ("p", "Li Ming"), ("p", "尾段")])

    def test_double_backslash_is_not_a_hard_break(self):
        # 正文真要写反斜杠时用两个，不该被当成强制换行
        self.assertEqual(gw.parse_body("结尾是双反斜杠\\\\\n下一行\n\n尾段\n"),
                         [("p", "结尾是双反斜杠\\\\下一行"), ("p", "尾段")])


class TextUtilTest(unittest.TestCase):
    def test_width_chars(self):
        self.assertEqual(gw.text_width_chars("××市"), 3.0)
        self.assertEqual(gw.text_width_chars("AB"), 1.0)

    def test_valid_date(self):
        self.assertTrue(gw._valid_date("2025 年 10 月 25 日"))
        self.assertTrue(gw._valid_date("2025年1月5日"))
        self.assertFalse(gw._valid_date("2025-10-25"))
        self.assertFalse(gw._valid_date("2025 年 2 月 30 日"))

    def test_valid_date_leap_year(self):
        self.assertTrue(gw._valid_date("2024 年 2 月 29 日"))
        self.assertFalse(gw._valid_date("2025 年 2 月 29 日"))

    def test_valid_date_out_of_range(self):
        self.assertFalse(gw._valid_date("2025 年 13 月 1 日"))
        self.assertFalse(gw._valid_date("2025 年 0 月 1 日"))
        self.assertFalse(gw._valid_date("2025 年 1 月 0 日"))

    def test_strip_attach_prefix(self):
        self.assertEqual(gw.strip_attach_prefix("附件：1. 某某表"), "1. 某某表")
        self.assertEqual(gw.strip_attach_prefix("2. 别的"), "2. 别的")


class ValidateArgsTest(unittest.TestCase):
    def test_ok(self):
        errors, notices = gw.validate_args(_args(title="正常标题"))
        self.assertEqual((errors, notices), ([], []))

    def test_out_must_be_docx(self):
        errors, _ = gw.validate_args(_args(out="a.txt"))
        self.assertTrue(any(".docx" in e for e in errors))

    def test_blank_title(self):
        errors, _ = gw.validate_args(_args(title="   "))
        self.assertTrue(any("--title" in e for e in errors))

    def test_to_needs_fullwidth_colon(self):
        errors, _ = gw.validate_args(_args(to="各处室"))
        self.assertTrue(any("冒号" in e for e in errors))
        errors, _ = gw.validate_args(_args(to="各处室："))
        self.assertEqual(errors, [])

    def test_blank_to_is_error(self):
        errors, _ = gw.validate_args(_args(to="   "))
        self.assertTrue(any("--to" in e for e in errors))
        self.assertEqual(gw.validate_args(_args(to=None))[0], [])

    def test_bad_date(self):
        errors, _ = gw.validate_args(_args(date="2025-10-25"))
        self.assertTrue(any("--date" in e for e in errors))

    def test_attach_prefix_is_notice_and_stripped_later(self):
        errors, notices = gw.validate_args(_args(attach=["附件：1. 某某表"]))
        self.assertEqual(errors, [])
        self.assertTrue(notices)


class ExitCodeTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def _main(self, argv, missing=([], None)):
        with mock.patch.object(gw, "find_missing_fonts", return_value=missing):
            return gw.main(argv)

    def test_check_fonts_ok(self):
        self.assertEqual(self._main(["--check-fonts"]), 0)

    def test_check_fonts_missing_returns_2(self):
        self.assertEqual(self._main(["--check-fonts"], ([("黑体", "一级标题")], None)), 2)

    def test_check_fonts_detect_error_returns_5(self):
        self.assertEqual(self._main(["--check-fonts"], ([("黑体", "一级标题")], "模拟检测失败")), 5)

    def test_generate_without_force_when_detect_error_returns_5(self):
        code = self._main(["--out", os.path.join(self.tmp, "x.docx"), "--title", "t"],
                          ([("黑体", "一级标题")], "模拟检测失败"))
        self.assertEqual(code, 5)

    def test_blank_title_returns_1(self):
        self.assertEqual(self._main(["--out", os.path.join(self.tmp, "x.docx"), "--title", "   "]), 1)

    def test_to_without_colon_returns_1(self):
        self.assertEqual(
            self._main(["--out", os.path.join(self.tmp, "x.docx"), "--title", "t", "--to", "各处室"]), 1)

    def test_empty_body_file_exits_1(self):
        body = os.path.join(self.tmp, "empty.txt")
        open(body, "w", encoding="utf-8").close()
        with self.assertRaises(SystemExit) as ctx:
            self._main(["--out", os.path.join(self.tmp, "x.docx"), "--title", "t", "--body", body])
        self.assertEqual(ctx.exception.code, 1)

    def test_non_utf8_body_exits_1_without_traceback(self):
        body = os.path.join(self.tmp, "gbk.txt")
        with open(body, "wb") as fh:
            fh.write("正文".encode("gbk"))
        with contextlib.redirect_stderr(io.StringIO()) as err, \
                self.assertRaises(SystemExit) as ctx:
            self._main(["--out", os.path.join(self.tmp, "x.docx"), "--title", "t", "--body", body])
        self.assertEqual(ctx.exception.code, 1)
        self.assertIn("UTF-8", err.getvalue())
        self.assertNotIn(self.tmp, err.getvalue())

    def test_page_problem_returns_4(self):
        def fake_build(args):
            def fake_save(path):
                with open(path, "wb") as fh:
                    fh.write(b"x")

            doc = mock.Mock()
            doc.save = fake_save
            return doc, [], ["页码炸了"]

        out = os.path.join(self.tmp, "x.docx")
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                mock.patch.object(gw, "build_document", side_effect=fake_build):
            code = gw.main(["--out", out, "--title", "t"])
        self.assertEqual(code, 4)
        # 退出码 4 的语义是「文档已生成，只是页码没写进去」，文件必须真的在
        self.assertTrue(os.path.isfile(out), "返回 4 但目标文件不存在")

    def test_unexpected_save_error_cleans_temp_and_reraises(self):
        def fake_build(args):
            def boom(path):
                with open(path, "wb") as fh:
                    fh.write(b"x")
                raise ValueError("内部炸了")

            doc = mock.Mock()
            doc.save = boom
            return doc, [], []

        out = os.path.join(self.tmp, "x.docx")
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                mock.patch.object(gw, "build_document", side_effect=fake_build):
            with self.assertRaises(ValueError):
                gw.main(["--out", out, "--title", "t"])
        # 非 OSError 的意外：临时文件照样清掉，异常原样抛出（不被吞）
        self.assertEqual([n for n in os.listdir(self.tmp) if n.endswith(".tmp")], [])
        self.assertFalse(os.path.exists(out))

    def test_validate_error_hides_absolute_path(self):
        secret = os.path.join(os.sep, "Users", "someone", "没有这个正文.txt")
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                contextlib.redirect_stderr(io.StringIO()) as err:
            code = gw.main(["--out", os.path.join(self.tmp, "x.docx"),
                            "--title", "t", "--body", secret])
        self.assertEqual(code, 1)
        self.assertIn("正文文件不存在", err.getvalue())
        self.assertIn("没有这个正文.txt", err.getvalue())
        self.assertNotIn("someone", err.getvalue())

    def test_out_path_error_hides_absolute_path(self):
        # --out 的报错同样要脱敏：这条分支此前没有测试盯着
        secret = os.path.join(os.sep, "Users", "someone", "错误的文件名.txt")
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                contextlib.redirect_stderr(io.StringIO()) as err:
            code = gw.main(["--out", secret, "--title", "t"])
        self.assertEqual(code, 1)
        self.assertIn("--out 必须以 .docx 结尾", err.getvalue())
        self.assertIn("错误的文件名.txt", err.getvalue())
        self.assertNotIn("someone", err.getvalue())

    def test_unexpected_save_error_with_cleanup_failure_keeps_original(self):
        # 保存抛 ValueError、清理又抛非 OSError：抛出来的必须还是原始 Exception
        def fake_build(args):
            def boom(path):
                with open(path, "wb") as fh:
                    fh.write(b"x")
                raise ValueError("内部炸了")

            doc = mock.Mock()
            doc.save = boom
            return doc, [], []

        out = os.path.join(self.tmp, "x.docx")
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                mock.patch.object(gw, "build_document", side_effect=fake_build), \
                mock.patch("os.unlink", side_effect=RuntimeError("清理也炸了")), \
                contextlib.redirect_stderr(io.StringIO()) as err:
            with self.assertRaises(ValueError):
                gw.main(["--out", out, "--title", "t"])
        text = err.getvalue()
        self.assertIn("临时文件没清掉", text)
        self.assertIn("RuntimeError", text)
        self.assertNotIn(self.tmp, text)
        self.assertFalse(os.path.exists(out))

    def test_save_error_returns_1_cleans_temp_and_hides_path(self):
        def fake_build(args):
            def boom(path):
                with open(path, "wb") as fh:
                    fh.write(b"x")
                raise OSError(28, "No space left on device")

            doc = mock.Mock()
            doc.save = boom
            return doc, [], []

        out = os.path.join(self.tmp, "x.docx")
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                mock.patch.object(gw, "build_document", side_effect=fake_build), \
                contextlib.redirect_stderr(io.StringIO()) as err:
            code = gw.main(["--out", out, "--title", "t"])
        self.assertEqual(code, 1)
        self.assertFalse(os.path.exists(out))
        self.assertEqual([n for n in os.listdir(self.tmp) if n.endswith(".tmp")], [])
        self.assertIn("No space left on device", err.getvalue())
        self.assertNotIn(self.tmp, err.getvalue())

    def test_out_pointing_to_directory_returns_1(self):
        target = os.path.join(self.tmp, "x.docx")
        os.makedirs(target)
        code = self._main(["--out", target, "--title", "t"])
        self.assertEqual(code, 1)

    def test_cleanup_failure_reports_and_returns_1(self):
        # 保存失败之后连删临时文件都失败：退出码仍是 1，并且必须提示残留
        def fake_build(args):
            def boom(path):
                with open(path, "wb") as fh:
                    fh.write(b"x")
                raise OSError(28, "No space left on device")

            doc = mock.Mock()
            doc.save = boom
            return doc, [], []

        out = os.path.join(self.tmp, "x.docx")
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                mock.patch.object(gw, "build_document", side_effect=fake_build), \
                mock.patch("os.unlink", side_effect=OSError(13, "Permission denied")), \
                contextlib.redirect_stderr(io.StringIO()) as err:
            code = gw.main(["--out", out, "--title", "t"])
        text = err.getvalue()
        self.assertEqual(code, 1)
        self.assertIn("临时文件没清掉", text)
        self.assertIn("Permission denied", text)
        self.assertNotIn(self.tmp, text)
        self.assertFalse(os.path.exists(out))

    def test_no_absolute_path_across_stdout_and_stderr(self):
        # 成功与失败两条通道都要脱敏：检测不可用时的提醒走 stderr，也要干净
        out = os.path.join(self.tmp, "out.docx")
        with mock.patch.object(gw, "find_missing_fonts",
                               return_value=([], "模拟：读字体目录失败：拒绝访问")), \
                contextlib.redirect_stdout(io.StringIO()) as out_buf, \
                contextlib.redirect_stderr(io.StringIO()) as err_buf:
            code = gw.main(["--out", out, "--title", "t", "--force"])
        self.assertEqual(code, 0)
        for buf in (out_buf.getvalue(), err_buf.getvalue()):
            self.assertNotIn(self.tmp, buf)
        self.assertIn("out.docx", out_buf.getvalue())
        self.assertIn("字体检测不可用", err_buf.getvalue())


class GenerateTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def _generate(self, body_text=None, **extra):
        argv = ["--out", os.path.join(self.tmp, "out.docx"), "--title", "关于开展××工作的通知",
                "--to", "各处室：", "--sign", "××市人民政府办公室",
                "--date", "2025 年 10 月 25 日",
                "--attach", "1. ××××表", "--attach", "2. ××××清单",
                "--note", "此件公开发布"]
        if body_text is not None:
            path = os.path.join(self.tmp, "body.txt")
            with open(path, "w", encoding="utf-8") as fh:
                fh.write(body_text)
            argv += ["--body", path]
        for key, value in extra.items():
            if value is True:
                argv.append(key)
            else:
                argv += [key, value]
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                contextlib.redirect_stdout(io.StringIO()):
            code = gw.main(argv)
        self.assertEqual(code, 0)
        return Document(os.path.join(self.tmp, "out.docx"))

    def test_page_setup(self):
        doc = self._generate()
        s = doc.sections[0]
        self.assertAlmostEqual(s.top_margin.mm, 37, places=1)
        self.assertAlmostEqual(s.bottom_margin.mm, 35, places=1)
        self.assertAlmostEqual(s.left_margin.mm, 28, places=1)
        self.assertAlmostEqual(s.right_margin.mm, 26, places=1)
        self.assertAlmostEqual(s.footer_distance.mm, 28, places=1)

    def test_title_font_and_centering(self):
        doc = self._generate()
        title = [p for p in doc.paragraphs if p.text.strip()][0]
        self.assertEqual(title.text, "关于开展××工作的通知")
        run = title.runs[0]
        self.assertEqual(run.font.size.pt, 22)
        rfonts = run._element.rPr.rFonts
        self.assertEqual(rfonts.get(qn("w:eastAsia")), "方正小标宋简体")
        self.assertEqual(title.alignment, WD_ALIGN_PARAGRAPH.CENTER)

    def test_body_font_and_indent(self):
        doc = self._generate("正文第一段。\n\n正文第二段。\n")
        body = [p for p in doc.paragraphs if p.text.startswith("正文第一段")][0]
        run = body.runs[0]
        self.assertEqual(run.font.size.pt, 16)
        self.assertEqual(run._element.rPr.rFonts.get(qn("w:eastAsia")), "仿宋_GB2312")
        ind = body._p.pPr.find(qn("w:ind"))
        self.assertEqual(ind.get(qn("w:firstLineChars")), "200")
        self.assertAlmostEqual(body.paragraph_format.line_spacing.pt, 29, places=1)

    def test_western_runs_use_times_new_roman(self):
        """成品文档里，凡是带英文数字的 run，ascii/hAnsi 都必须是 Times New Roman。"""
        doc = self._generate("第 3 条共 8 项，DeepSeek Harness 是工具名。\n")
        checked = 0
        for par in doc.paragraphs:
            if not par.runs or not par.text.strip():
                continue
            rfonts = par.runs[0]._element.rPr.rFonts
            self.assertEqual(rfonts.get(qn("w:ascii")), "Times New Roman", par.text)
            self.assertEqual(rfonts.get(qn("w:hAnsi")), "Times New Roman", par.text)
            checked += 1
        self.assertGreater(checked, 0)

    def test_heading_fonts(self):
        doc = self._generate("# 一、总体要求\n\n## （一）主要任务\n\n### 1. 第一项\n\n#### （1）细节\n")
        fonts = {p.text: p.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia"))
                 for p in doc.paragraphs if p.text.strip()}
        self.assertEqual(fonts["一、总体要求"], "黑体")
        self.assertEqual(fonts["（一）主要任务"], "楷体_GB2312")
        self.assertEqual(fonts["1. 第一项"], "仿宋_GB2312")
        self.assertEqual(fonts["（1）细节"], "仿宋_GB2312")

    def test_attachment_hanging_indent(self):
        doc = self._generate()
        first = [p for p in doc.paragraphs if p.text == "附件：1. ××××表"][0]
        second = [p for p in doc.paragraphs if p.text == "2. ××××清单"][0]
        ind1 = first._p.pPr.find(qn("w:ind"))
        self.assertEqual(ind1.get(qn("w:hangingChars")), "300")
        self.assertAlmostEqual(first.paragraph_format.left_indent.pt, 80, places=1)
        self.assertAlmostEqual(first.paragraph_format.first_line_indent.pt, -48, places=1)
        # 第 2 条与第 1 条的「1.」对齐（左空 2 + 「附件：」3 = 5 字），另加 1.5 字续行悬挂
        ind2 = second._p.pPr.find(qn("w:ind"))
        self.assertEqual(ind2.get(qn("w:hangingChars")), "150")
        self.assertAlmostEqual(second.paragraph_format.left_indent.pt, 104, places=1)
        self.assertAlmostEqual(second.paragraph_format.first_line_indent.pt, -24, places=1)

    def test_attachment_skips_blank_and_keeps_alignment(self):
        # 传了空串时，实际第一条仍要带「附件：」，不能把第 2 条顶到「附件：」位置
        args = _args(attach=["  ", "1. 表"], sign=None, note=None)
        doc, _blocks, _problems = gw.build_document(args)
        attach = [p for p in doc.paragraphs if p.text.startswith("附件：")]
        self.assertEqual([p.text for p in attach], ["附件：1. 表"])

    def test_all_blank_attach_leaves_no_extra_blank_par(self):
        # 整组附件都是空白：不该在正文后多留一个空段
        base = dict(sign=None, note=None, date=None, to=None, body=None)
        without = gw.build_document(_args(attach=None, **base))[0]
        blanks = gw.build_document(_args(attach=["   ", "\t"], **base))[0]
        self.assertEqual([p.text for p in without.paragraphs],
                         [p.text for p in blanks.paragraphs])

    def test_page_number_field(self):
        doc = self._generate()
        xml = doc.sections[0].footer.paragraphs[0]._p.xml
        self.assertIn("PAGE", xml)
        self.assertIn("— ", doc.sections[0].footer.paragraphs[0].text)

    def test_page_number_layout(self):
        doc = self._generate()
        odd = doc.sections[0].footer.paragraphs[0]
        even = doc.sections[0].even_page_footer.paragraphs[0]
        self.assertEqual(odd.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(even.alignment, WD_ALIGN_PARAGRAPH.LEFT)
        for par in (odd, even):
            self.assertIn("PAGE", par._p.xml)
            self.assertEqual(par.runs[0].font.size.pt, 14)
            self.assertEqual(par.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")), "宋体")
        self.assertTrue(doc.settings.odd_and_even_pages_header_footer)

    def test_hide_first_page_number(self):
        doc = self._generate(**{"--hide-first-page-number": True})
        self.assertTrue(doc.sections[0].different_first_page_header_footer)

    def test_note_wrapped_in_parentheses(self):
        doc = self._generate()
        note = [p for p in doc.paragraphs if "此件公开发布" in p.text][0]
        self.assertEqual(note.text, "（此件公开发布）")

    def test_no_absolute_path_in_log(self):
        buf = io.StringIO()
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                contextlib.redirect_stdout(buf):
            gw.main(["--out", os.path.join(self.tmp, "out.docx"), "--title", "t"])
        self.assertNotIn(self.tmp, buf.getvalue())
        self.assertIn("out.docx", buf.getvalue())

    def test_atomic_write_leaves_no_tmp(self):
        self._generate()
        self.assertEqual([f for f in os.listdir(self.tmp) if f.endswith(".tmp")], [])

    def test_output_dir_is_created(self):
        deep = os.path.join(self.tmp, "a", "b", "out.docx")
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                contextlib.redirect_stdout(io.StringIO()):
            code = gw.main(["--out", deep, "--title", "t"])
        self.assertEqual(code, 0)
        self.assertTrue(os.path.isfile(deep))


class FontPartialDetectionTest(unittest.TestCase):
    """「部分来源读不到」和「一条都读不到」要区分开：前者只提醒，后者才抬到 5。"""

    def test_partial_failure_warns_but_stays_below_exit_5(self):
        def fake_installed():
            # 注册表读失败、字体目录正常：结论不可信，但还没到「检测不可用」
            return gw.FontScan(set(), {"SIMFANG", "SIMKAI"}, [],
                               ["读注册表 HKLM（本机）失败：拒绝访问"])

        with mock.patch.object(gw, "installed_fonts", side_effect=fake_installed), \
                contextlib.redirect_stderr(io.StringIO()) as err:
            missing, detect_error = gw.find_missing_fonts()
        self.assertIsNone(detect_error)
        self.assertTrue(missing)  # 只靠文件名兜底，结论可能少算但不会漏到空
        self.assertIn("字体检测不完整", err.getvalue())

    def test_nothing_read_at_all_is_detect_error(self):
        with mock.patch.object(
                gw, "installed_fonts",
                return_value=gw.FontScan(set(), set(), [], ["读注册表失败：拒绝访问"])):
            missing, detect_error = gw.find_missing_fonts()
        self.assertIsNotNone(detect_error)
        self.assertEqual(len(missing), len(gw.REQUIRED_FONTS))

    def test_name_table_rescues_font_with_odd_filename(self):
        """方正小标宋简体装成 FZXBSJW.TTF，注册表和文件名都认不出，靠 name 表翻案。"""
        scan = gw.FontScan(set(), {"FZXBSJW"}, ["/fake/FZXBSJW.TTF"], [])
        with mock.patch.object(gw, "installed_fonts", return_value=scan), \
                mock.patch.object(gw, "_font_names_from_file",
                                  return_value=frozenset({"方正小标宋简体"})):
            missing, detect_error = gw.find_missing_fonts()
        self.assertIsNone(detect_error)
        self.assertNotIn("方正小标宋简体", [font for font, _ in missing])

    def test_no_parse_when_scan_already_sufficient(self):
        """名字全对得上时不该再去翻字体文件——解析 name 表是慢路径。"""
        scan = gw.FontScan({font for font, _ in gw.REQUIRED_FONTS}, set(), [], [])
        with mock.patch.object(gw, "installed_fonts", return_value=scan), \
                mock.patch.object(gw, "_font_names_from_file") as parse:
            missing, _ = gw.find_missing_fonts()
        self.assertEqual(missing, [])
        parse.assert_not_called()


class FontNameTableTest(unittest.TestCase):
    """_parse_name_table / _norm_font_name / _gdi_font_families 这类底层解析。

    只造最小二进制，不依赖本机装了什么字体，CI 上也能跑。
    """

    @staticmethod
    def _build_sfnt(names):
        """按 name 表结构拼一个最小 sfnt。

        names 是 [(platform, encoding, language, nameID, text)]；text 由调用方
        决定该用 utf-16-be 还是 mac_roman 编码。
        """
        records = b""
        strings = b""
        for platform, encoding, language, name_id, blob in names:
            records += struct.pack(">HHHHHH", platform, encoding, language,
                                   name_id, len(blob), len(strings))
            strings += blob
        header = struct.pack(">HHH", 0, len(names), 6 + len(records))
        return header + records + strings

    def test_reads_utf16_chinese_name(self):
        blob = "方正小标宋简体".encode("utf-16-be")
        table = self._build_sfnt([(3, 1, 0x804, 1, blob), (3, 1, 0x804, 4, blob)])
        self.assertEqual(gw._parse_name_table(table), {"方正小标宋简体"})

    def test_reads_latin_name_from_mac_roman(self):
        table = self._build_sfnt([(1, 0, 0, 1, b"Times New Roman")])
        self.assertEqual(gw._parse_name_table(table), {"Times New Roman"})

    def test_truncated_table_returns_empty(self):
        blob = "黑体".encode("utf-16-be")
        table = self._build_sfnt([(3, 1, 0x804, 1, blob)])
        self.assertEqual(gw._parse_name_table(table[:-6]), set())

    def test_normalisation_ignores_space_case_and_suffix(self):
        self.assertEqual(gw._norm_font_name("Times New Roman"),
                         gw._norm_font_name("times_new-roman"))
        self.assertEqual(gw._norm_font_name("黑体"), gw._norm_font_name("黑 体"))

    def test_gdi_off_on_non_windows(self):
        with mock.patch.object(gw.sys, "platform", "linux"):
            self.assertEqual(gw._gdi_font_families(), set())


class InstalledFontsSourceTest(unittest.TestCase):
    """不桩 installed_fonts 本身，改为桩它下面真正的来源（winreg / listdir / GDI）。"""

    def _fake_winreg(self, openkey_error):
        fake = mock.MagicMock()
        fake.HKEY_LOCAL_MACHINE = "HKLM"
        fake.HKEY_CURRENT_USER = "HKCU"
        # 三套视图的位标志要是整数，否则 base_flags | ... 出来的是 MagicMock
        fake.KEY_READ = 0x20019
        fake.KEY_WOW64_64KEY = 0x0100
        fake.KEY_WOW64_32KEY = 0x0200
        fake.OpenKey.side_effect = openkey_error
        return fake

    def _installed_with(self, fake_reg, listdir):
        """在打桩 winreg / listdir / GDI 的前提下跑一次 installed_fonts。

        GDI 必须一起桩掉：本机真实存在字体，不桩的话「全部失败」的用例永远
        撞上非空结果。
        """
        with mock.patch.object(gw.sys, "platform", "win32"), \
                mock.patch.dict(sys.modules, {"winreg": fake_reg}), \
                mock.patch.object(gw.os, "listdir", side_effect=listdir), \
                mock.patch.object(gw, "_gdi_font_families", return_value=set()):
            return gw.installed_fonts()

    def test_registry_error_but_font_dir_readable(self):
        fake = self._fake_winreg(OSError(5, "拒绝访问"))
        scan = self._installed_with(fake, lambda _d: ["simfang.ttf"])
        self.assertEqual(scan.names, set())
        self.assertIn("SIMFANG", scan.files)
        self.assertTrue(any("读注册表" in e for e in scan.errors))
        self.assertFalse(any("读系统字体目录" in e for e in scan.errors))

    def test_registry_ok_fills_names(self):
        fake = mock.MagicMock()
        fake.HKEY_LOCAL_MACHINE = "HKLM"
        fake.HKEY_CURRENT_USER = "HKCU"
        fake.KEY_READ = 0x20019
        fake.KEY_WOW64_64KEY = 0x0100
        fake.KEY_WOW64_32KEY = 0x0200
        key = mock.MagicMock()
        fake.OpenKey.return_value.__enter__.return_value = key
        fake.QueryInfoKey.return_value = (0, 1, 0)
        fake.EnumValue.return_value = ("黑体 (TrueType)", "simhei.ttf", 1)
        scan = self._installed_with(fake, lambda _d: [])
        self.assertIn("黑体", scan.names)
        self.assertIn("SIMHEI", scan.files)
        self.assertEqual(scan.errors, [])

    def test_all_sources_fail_is_detect_error(self):
        fake = self._fake_winreg(OSError(5, "拒绝访问"))
        with mock.patch.object(gw.sys, "platform", "win32"), \
                mock.patch.dict(sys.modules, {"winreg": fake}), \
                mock.patch.object(gw.os, "listdir", side_effect=OSError(3, "路径不存在")), \
                mock.patch.object(gw, "_gdi_font_families", return_value=set()):
            missing, detect_error = gw.find_missing_fonts()
        self.assertIsNotNone(detect_error)
        self.assertEqual(len(missing), len(gw.REQUIRED_FONTS))

    def test_user_font_dir_is_scanned(self):
        """用户自装字体在 %LOCALAPPDATA%\\Microsoft\\Windows\\Fonts，不在系统目录。"""
        fake = self._fake_winreg(OSError(5, "拒绝访问"))
        user_dir = os.path.join("C:\\Users\\x\\AppData\\Local",
                                "Microsoft", "Windows", "Fonts")

        def listdir(directory):
            if directory == user_dir:
                return ["FZXBSJW.TTF"]
            raise OSError(3, "路径不存在")

        with mock.patch.dict(os.environ, {"LOCALAPPDATA": "C:\\Users\\x\\AppData\\Local"}):
            scan = self._installed_with(fake, listdir)
        self.assertIn("FZXBSJW", scan.files)
        self.assertTrue(any(p.endswith("FZXBSJW.TTF") for p in scan.paths))

    def test_non_windows_reports_error_and_empty(self):
        with mock.patch.object(gw.sys, "platform", "linux"):
            scan = gw.installed_fonts()
        self.assertEqual(scan.names, set())
        self.assertEqual(scan.files, set())
        self.assertTrue(scan.errors)


class WesternFontSplitTest(unittest.TestCase):
    """英文和数字走 Times New Roman，中文仍走中文字体——靠 rFonts 四个属性分流。"""

    def _set_run_font(self):
        return gw.make_helpers(*gw._import_docx())[0]

    def test_set_run_font_splits_east_asian_and_western(self):
        doc = Document()
        run = doc.add_paragraph().add_run("甲乙 ABC 123")
        self._set_run_font()(run, "仿宋_GB2312", gw.SIZE_BODY)
        rfonts = run._element.rPr.rFonts
        self.assertEqual(rfonts.get(qn("w:eastAsia")), "仿宋_GB2312")
        self.assertEqual(rfonts.get(qn("w:ascii")), "Times New Roman")
        self.assertEqual(rfonts.get(qn("w:hAnsi")), "Times New Roman")
        self.assertEqual(rfonts.get(qn("w:cs")), "Times New Roman")

    def test_times_new_roman_is_required(self):
        self.assertIn(gw.FONT_WESTERN, [font for font, _ in gw.REQUIRED_FONTS])


class HeadingWarningsTest(unittest.TestCase):
    """层级标记和中文编号不配套时要提醒——写错了字体全错，肉眼很难发现。"""

    def test_matching_numbering_is_silent(self):
        blocks = [("h1", "一、总则"), ("h2", "（一）范围"),
                  ("h3", "1. 适用"), ("h4", "（1）细则")]
        self.assertEqual(gw.heading_warnings(blocks), [])

    def test_unnumbered_heading_is_silent(self):
        # 不是所有标题都带编号，「使用说明」这种不该被提醒
        self.assertEqual(gw.heading_warnings([("h1", "使用说明"), ("h2", "总则")]), [])

    def test_h2_with_level_one_numbering_warns(self):
        warnings = gw.heading_warnings([("h2", "一、大作文")])
        self.assertEqual(len(warnings), 1)
        self.assertIn("##", warnings[0])
        self.assertIn("（一）", warnings[0])

    def test_h3_with_paren_numbering_warns(self):
        warnings = gw.heading_warnings([("h3", "（一）这里该用数字")])
        self.assertEqual(len(warnings), 1)
        self.assertIn("###", warnings[0])

    def test_h4_with_arabic_ordinal_warns(self):
        warnings = gw.heading_warnings([("h4", "1. 该用括号数字")])
        self.assertEqual(len(warnings), 1)
        self.assertIn("（1）", warnings[0])

    def test_paragraphs_are_never_checked(self):
        self.assertEqual(gw.heading_warnings([("p", "一、正文里的编号")]), [])


class VerifyOutputTest(unittest.TestCase):
    """--verify 自检：合规文档放行，字体写错的要抓出来。"""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.out = os.path.join(self.tmp, "out.docx")

    def _generate(self, body_text):
        path = os.path.join(self.tmp, "body.txt")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write(body_text)
        argv = ["--out", self.out, "--title", "关于开展××工作的通知", "--body", path]
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(gw.main(argv), 0)

    def test_clean_document_passes(self):
        self._generate("# 一、总则\n\n正文第 1 段，含 English 与数字 2025。\n")
        self.assertEqual(gw.verify_output(self.out), [])

    def test_verify_flag_returns_0_on_clean_document(self):
        path = os.path.join(self.tmp, "body.txt")
        with open(path, "w", encoding="utf-8") as fh:
            fh.write("正文第 1 段。\n")
        argv = ["--out", self.out, "--title", "关于开展××工作的通知",
                "--body", path, "--verify"]
        with mock.patch.object(gw, "find_missing_fonts", return_value=([], None)), \
                contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(gw.main(argv), 0)

    def test_grabs_western_font_that_is_not_times_new_roman(self):
        # 手工把某个 run 的西文字体改掉，模拟旧版本脚本生成的文件
        self._generate("正文第 1 段。\n")
        doc = Document(self.out)
        run = [p for p in doc.paragraphs if p.text.strip()][0].runs[0]
        run._element.rPr.rFonts.set(qn("w:ascii"), "仿宋_GB2312")
        doc.save(self.out)
        problems = gw.verify_output(self.out)
        self.assertTrue(any("Times New Roman" in p for p in problems), problems)


if __name__ == "__main__":
    unittest.main()
