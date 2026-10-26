#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""公文版式 Word 生成器（GB/T 9704-2012 版式部分，无红头）。

用法示例：
  # 只查字体
  python gongwen_docx.py --check-fonts

  # 生成文档
  python gongwen_docx.py --out 输出.docx \
      --title "关于××××的通知" --to "各处室：" --body 正文.txt \
      --sign "××××办公室" --date "2025年10月25日" \
      --attach "1. ××××表" --note "（此件公开发布）"

正文文件（--body）的层级标记：
  # 一级标题      → 3 号黑体        「一、」
  ## 二级标题      → 3 号楷体_GB2312 「（一）」
  ### 三级标题     → 3 号仿宋        「1.」
  #### 四级标题    → 3 号仿宋        「（1）」
  其余非空行       → 正文段（3 号仿宋，首行缩进 2 字）

正文行首真的要以 # 开头时，写成 \\# —— 反斜杠会被吃掉，井号原样保留。

分段规则：空行分段。连续的非空普通行合并成同一个正文段（段内中文直接相连，
英文/数字相接处补一个半角空格）；空行本身不产生空段。整份文件一个空行都没有
时，退回「一行一段」，免得整篇正文被合并成一大段。

层级标记只换字体，不自动加编号——「一、」「（一）」「1.」要自己写进文本。

中文字体按 GB/T 9704-2012 走；英文和阿拉伯数字一律 Times New Roman（中文字符用
w:eastAsia，西文用 w:ascii / w:hAnsi，Word 自己分流）。

字体检测只支持 Windows，但会同时走四条路：注册表（HKLM + HKCU，默认/64 位/32 位
三套视图）、系统与用户两个字体目录、字体文件 name 表里自报的中文名、GDI 字体族
枚举。用户自装的字体常常不在 C:\\Windows\\Fonts、不进注册表、文件名也不是中文名
（方正小标宋简体是 FZXBSJW.TTF），只看其中一条会误判成没装。非 Windows，或一条
字体信息都读不出来时，脚本不猜：按「无法确认」处理，退出码 5，加 --force 可放行
继续。

字体缺失时文件名仍按规范写入，脚本不做替换；--force / --allow-missing-fonts
只负责放行，实际显示由打开文档的机器决定。

退出码：0 正常；1 参数、读取或写入错误；2 缺字体未获允许；3 缺依赖；
        4 文档已生成但页码设置失败；5 字体检测不可用，未获允许；
        6 文档已生成，但 --verify 自检发现问题。
"""

from __future__ import annotations

import argparse
import datetime
import os
import re
import struct
import sys
import tempfile
from typing import NamedTuple


def _setup_console() -> None:
    """把控制台输出切成 UTF-8。

    Windows 上默认编码可能是 cp1252 / cp936，打印中文直接抛 UnicodeEncodeError，
    会让已经生成好的文档卡在最后一步。errors="replace" 保证不崩。

    只在 main() 里调用：这个脚本也会被测试或别的程序 import，import 时不该顺手
    改掉宿主的 stdout/stderr。
    """
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

# ---------------------------------------------------------------- 版式常量

MM_TOP = 37          # 页边距 上
MM_BOTTOM = 35       # 页边距 下
MM_LEFT = 28         # 页边距 左
MM_RIGHT = 26        # 页边距 右
LINE_PITCH_PT = 29   # 固定行距（225mm ÷ 22 行 ≈ 29 磅）
LINE_PITCH_TWIPS = 580  # 同上，换成 twips 写进 docGrid
INDENT_CHARS = 2     # 首行缩进
SIZE_TITLE = 22      # 2 号
SIZE_BODY = 16       # 3 号
SIZE_PAGE = 14       # 4 号

FONT_TITLE = "方正小标宋简体"
FONT_BODY = "仿宋_GB2312"
FONT_H1 = "黑体"
FONT_H2 = "楷体_GB2312"
FONT_H3 = "仿宋_GB2312"
FONT_PAGE = "宋体"
# 英文和阿拉伯数字一律走这款；中文字符仍用上面各自的字体。
FONT_WESTERN = "Times New Roman"

REQUIRED_FONTS = [
    (FONT_TITLE, "标题"),
    (FONT_BODY, "正文"),
    (FONT_H1, "一级标题"),
    (FONT_H2, "二级标题"),
    (FONT_PAGE, "页码"),
    (FONT_WESTERN, "英文与数字"),
]

# 最后一道兜底：注册表、字体目录、name 表全都失手时，靠文件名主干认字体
# （值统一大写、不含扩展名；字体自带的文件名五花八门，靠中文名根本对不上）
FONT_FILE_HINTS = {
    FONT_TITLE: {"FZXBSJW", "FZXIAOBIAOSONG-B05S", "FZXBSK", "FZXBSJW--GB1-0"},
    FONT_BODY: {"FZFSJW", "FZFSK", "SIMFANG", "FZFSJW--GB1-0"},
    FONT_H2: {"FZKTJW", "FZKATI", "SIMKAI", "FZKTJW--GB1-0", "KAIU"},
    FONT_H1: {"SIMHEI", "FZHTJW"},
    FONT_PAGE: {"SIMSUN", "SIMSUNB"},
    FONT_WESTERN: {"TIMES", "TIMESNEWROMAN", "TIMESNEWROMANPSMT"},
}

# 字体文件后缀：.ttc/.otc 是字体集合，一个文件里打包了多款字体
FONT_EXTS = (".ttf", ".otf", ".ttc", ".otc")


def _os_error_text(exc: OSError) -> str:
    """只取错误描述，不带路径——日志可能原样贴到 issue 里。"""
    return exc.strerror or exc.__class__.__name__


def _error_text(exc: BaseException) -> str:
    """清理阶段的错误描述。非 OSError 只留类名：异常消息里可能带本机路径。"""
    if isinstance(exc, OSError):
        return exc.strerror or exc.__class__.__name__
    return exc.__class__.__name__


_FONT_NAME_SUFFIXES = (
    " (TrueType)", " (OpenType)", " (All res)", " (PostScript)", " (Type 1)",
)


def _strip_font_suffix(name: str) -> str:
    """去掉注册表值名尾部的格式后缀：「黑体 (TrueType)」→「黑体」。"""
    name = name.strip()
    changed = True
    while changed:
        changed = False
        for suffix in _FONT_NAME_SUFFIXES:
            if name.endswith(suffix):
                name = name[: -len(suffix)].strip()
                changed = True
    return name


def _clean_font_name(text: str) -> str:
    """字体名清理：剔掉不可见控制符、掐掉首尾空白，过长的一律不要。"""
    text = "".join(ch for ch in text if ch.isprintable()).strip()
    if not text or len(text) > 128:
        return ""
    return text


def _norm_font_name(name: str) -> str:
    """比对用归一化：忽略空格、下划线、连字符和大小写。

    同一款字体在这种写法上五花八门（仿宋_GB2312 / 仿宋 GB2312 / 仿宋GB2312），
    归一化后落到同一个键，免得因为写法差异被判成没装。
    """
    return re.sub(r"[\s_\-]+", "", name).casefold()


# ---------------------------------------------------------------- 字体来源
#
# 只信一条路会翻车：注册表认名字但用户自装字体常常不进注册表；扫 C:\Windows
# \Fonts 认文件名但字体可能装在 %LOCALAPPDATA%\Microsoft\Windows\Fonts，而且
# 文件名还不一定是中文名（方正小标宋简体就是 FZXBSJW.TTF）。所以四条路都走，
# 任何一条认出来就算装了。

# 字体文件 name 表解析缓存：路径 → 名字集合。同一进程内一个文件只读一次。
_NAME_CACHE: dict[str, frozenset[str]] = {}


def _parse_name_table(blob: bytes) -> set[str]:
    """解析 sfnt 的 name 表，取出字体自报的中文名与英文名。

    name 表结构：format(2) count(2) stringOffset(2)，随后 count 条各 12 字节的
    nameRecord（platformID / encodingID / languageID / nameID / length / offset），
    字符串接在 stringOffset 之后。Windows 平台的字符串是 UTF-16BE——
    「方正小标宋简体」这种中文名就躺在里面，跟文件名长什么样完全无关。
    """
    if len(blob) < 6:
        return set()
    count, storage = struct.unpack(">HH", blob[2:6])
    if count > 4096 or storage > len(blob):
        return set()
    names: set[str] = set()
    for index in range(count):
        record = blob[6 + index * 12: 6 + (index + 1) * 12]
        if len(record) < 12:
            break
        (platform, encoding, _language,
         name_id, length, offset) = struct.unpack(">HHHHHH", record)
        # 只认 1 族名 / 4 全名 / 6 PostScript 名 / 16 排版族名，其余是版权说明之类
        if name_id not in (1, 4, 6, 16):
            continue
        start = storage + offset
        raw = blob[start:start + length]
        if len(raw) != length:
            continue
        text = ""
        if platform in (0, 3) and encoding in (1, 10):   # Unicode / Windows UCS-2
            text = raw.decode("utf-16-be", "replace")
        elif platform == 1 and encoding == 0:            # Mac Roman
            text = raw.decode("mac_roman", "replace")
        if not text:
            continue
        cleaned = _clean_font_name(text)
        if cleaned:
            names.add(cleaned)
    return names


def _read_sfnt_names(handle, base: int) -> set[str]:
    """从 handle 的 base 偏移处（一个 sfnt 头）读出 name 表里的名字。"""
    handle.seek(base)
    header = handle.read(12)
    if len(header) < 12:
        return set()
    num_tables = struct.unpack(">H", header[4:6])[0]
    if not 0 < num_tables <= 512:
        return set()
    records = handle.read(num_tables * 16)
    for index in range(min(num_tables, len(records) // 16)):
        record = records[index * 16:(index + 1) * 16]
        if record[:4] != b"name":
            continue
        offset, length = struct.unpack(">II", record[8:16])
        if not 0 < length <= 4 * 1024 * 1024:
            return set()
        handle.seek(offset)
        return _parse_name_table(handle.read(length))
    return set()


def _font_names_from_file(path: str) -> frozenset[str]:
    """直接读一个字体文件自报的名字；读不动就返回空集，绝不影响其它来源。"""
    cached = _NAME_CACHE.get(path)
    if cached is not None:
        return cached
    found: set[str] = set()
    try:
        with open(path, "rb") as handle:
            magic = handle.read(4)
            if magic == b"ttcf":
                # 字体集合：文件头后面跟一串各字体的偏移
                handle.seek(8)
                raw_count = handle.read(4)
                count = struct.unpack(">I", raw_count)[0] if len(raw_count) == 4 else 0
                count = min(count, 64)
                offsets = handle.read(count * 4)
                for index in range(count):
                    chunk = offsets[index * 4:(index + 1) * 4]
                    if len(chunk) < 4:
                        break
                    found |= _read_sfnt_names(handle, struct.unpack(">I", chunk)[0])
            elif magic in (b"\x00\x01\x00\x00", b"OTTO", b"true", b"ttff"):
                found = _read_sfnt_names(handle, 0)
    except (OSError, struct.error, ValueError):
        found = set()
    result = frozenset(found)
    _NAME_CACHE[path] = result
    return result


def _gdi_font_families() -> set[str]:
    """用 GDI 枚举本机字体族名。

    这条兜底既不看注册表也不看文件名，字体装在哪个目录都能看见；代价是只在
    Windows 上可用。调用失败就静默返回空集——一条兜底坏了不该拖垮整体结论。
    """
    if sys.platform != "win32":
        return set()
    try:
        import ctypes
        from ctypes import wintypes
    except ImportError:
        return set()

    class LOGFONTW(ctypes.Structure):
        _fields_ = [
            ("lfHeight", wintypes.LONG),
            ("lfWidth", wintypes.LONG),
            ("lfEscapement", wintypes.LONG),
            ("lfOrientation", wintypes.LONG),
            ("lfWeight", wintypes.LONG),
            ("lfItalic", wintypes.BYTE),
            ("lfUnderline", wintypes.BYTE),
            ("lfStrikeOut", wintypes.BYTE),
            ("lfCharSet", wintypes.BYTE),
            ("lfOutPrecision", wintypes.BYTE),
            ("lfClipPrecision", wintypes.BYTE),
            ("lfQuality", wintypes.BYTE),
            ("lfPitchAndFamily", wintypes.BYTE),
            ("lfFaceName", wintypes.WCHAR * 32),
        ]

    collected: set[str] = set()

    @ctypes.WINFUNCTYPE(ctypes.c_int, ctypes.POINTER(LOGFONTW),
                        ctypes.c_void_p, wintypes.DWORD, wintypes.LPARAM)
    def _collect(logfont, _metric, _font_type, _param):
        name = _clean_font_name(logfont.contents.lfFaceName)
        if name:
            collected.add(name)
        return 1

    try:
        gdi32 = ctypes.WinDLL("gdi32")
        gdi32.CreateCompatibleDC.restype = ctypes.c_void_p
        gdi32.CreateCompatibleDC.argtypes = [ctypes.c_void_p]
        gdi32.DeleteDC.argtypes = [ctypes.c_void_p]
        gdi32.EnumFontFamiliesExW.restype = ctypes.c_int
        gdi32.EnumFontFamiliesExW.argtypes = [
            ctypes.c_void_p, ctypes.POINTER(LOGFONTW),
            type(_collect), wintypes.LPARAM, wintypes.DWORD]
        hdc = gdi32.CreateCompatibleDC(None)
        if not hdc:
            return set()
        try:
            request = LOGFONTW()
            request.lfCharSet = 1  # DEFAULT_CHARSET：不限定字符集，全都要
            gdi32.EnumFontFamiliesExW(hdc, ctypes.byref(request), _collect, 0, 0)
        finally:
            gdi32.DeleteDC(hdc)
    except (OSError, AttributeError, ValueError):
        return set()
    return collected


class FontScan(NamedTuple):
    """一次字体扫描的全部结果。

    names   注册表和 GDI 报出来的字体名
    files   字体目录与注册表给出的文件名主干（大写、不含扩展名）
    paths   字体文件路径，给 name 表按需解析用
    errors  各来源的失败描述；空列表表示全部读成功

    解析 name 表比较慢，所以这里只收集路径，等 find_missing_fonts 发现确有
    缺口时才去翻——正常情况下一分钱开销都不花。
    """

    names: set[str]
    files: set[str]
    paths: list[str]
    errors: list[str]


def installed_fonts() -> FontScan:
    """扫本机字体：注册表（HKLM + HKCU，默认/64 位/32 位三套视图）+ 两个字体目录 + GDI。

    每条来源各自尽力、失败只记账不中断。只有一条都没读到，才算「检测不可用」。
    """
    names: set[str] = set()
    files: set[str] = set()
    paths: set[str] = set()
    errors: list[str] = []

    if sys.platform != "win32":
        errors.append("当前平台 %s 不是 Windows，读不到字体注册表、字体目录和 GDI 字体列表"
                      % sys.platform)
        return FontScan(names, files, [], errors)

    win_dir = os.environ.get("WINDIR", r"C:\Windows")

    try:
        import winreg
    except ImportError:
        winreg = None  # type: ignore
    if winreg is not None:
        sub = r"SOFTWARE\Microsoft\Windows NT\CurrentVersion\Fonts"
        base_flags = getattr(winreg, "KEY_READ", 0x20019)
        # 32 位 Python 会被重定向到 WOW6432Node，三套视图都试才不会漏
        views = [base_flags]
        if hasattr(winreg, "KEY_WOW64_64KEY"):
            views.append(base_flags | winreg.KEY_WOW64_64KEY)
        if hasattr(winreg, "KEY_WOW64_32KEY"):
            views.append(base_flags | winreg.KEY_WOW64_32KEY)
        for root, root_name in ((winreg.HKEY_LOCAL_MACHINE, "HKLM（本机）"),
                                (winreg.HKEY_CURRENT_USER, "HKCU（当前用户）")):
            read_any = False
            last_error: OSError | None = None
            for flags in views:
                try:
                    with winreg.OpenKey(root, sub, 0, flags) as key:
                        count = winreg.QueryInfoKey(key)[1]
                        for index in range(count):
                            name, value, _kind = winreg.EnumValue(key, index)
                            names.add(_strip_font_suffix(str(name)))
                            raw = str(value)
                            base = os.path.splitext(os.path.basename(raw))[0]
                            if base:
                                files.add(base.upper())
                            if os.path.splitext(raw)[1].lower() in FONT_EXTS:
                                paths.add(raw if os.path.isabs(raw)
                                          else os.path.join(win_dir, "Fonts", raw))
                    read_any = True
                except OSError as exc:
                    last_error = exc
            if not read_any:
                errors.append("读注册表 %s 失败：%s"
                              % (root_name,
                                 _os_error_text(last_error) if last_error
                                 else "没有可用的读取权限"))
    else:
        errors.append("当前 Python 没有 winreg 模块，读不到字体注册表")

    # 字体目录兜底：用户自装的字体常常不在 C:\Windows\Fonts，也不进注册表
    dirs = [("系统字体目录", os.path.join(win_dir, "Fonts"))]
    local_app = os.environ.get("LOCALAPPDATA", "")
    if local_app:
        dirs.append(("用户字体目录",
                     os.path.join(local_app, "Microsoft", "Windows", "Fonts")))
    for label, directory in dirs:
        try:
            for filename in os.listdir(directory):
                stem, ext = os.path.splitext(filename)
                if not stem:
                    continue
                files.add(stem.upper())
                if ext.lower() in FONT_EXTS:
                    paths.add(os.path.join(directory, filename))
        except OSError as exc:
            errors.append("读%s失败：%s" % (label, _os_error_text(exc)))

    names |= _gdi_font_families()
    return FontScan(names, files, sorted(paths), errors)


def _font_matches(font: str, scan: FontScan, parsed: set[str]) -> bool:
    """这款字体能不能算「装了」：名字、文件名、name 表三路任意一路对上就算。"""
    target = _norm_font_name(font)
    for reported in scan.names:
        if _norm_font_name(reported) == target:
            return True
    for reported in parsed:
        if _norm_font_name(reported) == target:
            return True
    if target.upper() in scan.files:
        return True
    if FONT_FILE_HINTS.get(font, set()) & scan.files:
        return True
    return False


def find_missing_fonts() -> tuple[list[tuple[str, str]], str | None]:
    """返回 (缺失字体, 检测不可用的原因)。

    原因是 None 表示检测生效、结论可信；非 None 表示一条字体信息都没读到，
    此时「缺失」只是全部必需字体，不代表真的没装。
    """
    scan = installed_fonts()
    if not scan.names and not scan.files:
        reason = "；".join(scan.errors) if scan.errors else "没有读到任何已安装字体"
        return [(font, role) for font, role in REQUIRED_FONTS], reason
    if scan.errors:
        # 部分来源读失败：结论可能少算，但还没到「一条都没读到」的程度，
        # 所以不把退出码抬到 5，只用一句提醒把不确定性说清楚。
        print("提醒：字体检测不完整（%s），缺失结论可能不准。" % "；".join(scan.errors),
              file=sys.stderr)

    missing = [item for item in REQUIRED_FONTS
               if not _font_matches(item[0], scan, set())]
    if not missing:
        return [], None

    # 还剩没认出来的：翻字体文件自报的名字。只在这一步付出解析开销。
    parsed: set[str] = set()
    for path in scan.paths:
        parsed |= _font_names_from_file(path)
    still_missing = [item for item in missing
                     if not _font_matches(item[0], scan, parsed)]
    return still_missing, None


# ---------------------------------------------------------------- docx 依赖

def _import_docx():
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt
    except ImportError:
        print("缺少依赖 python-docx。请先执行：pip install python-docx", file=sys.stderr)
        raise SystemExit(3)
    return (Document, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, OxmlElement, qn, Mm, Pt)


# ---------------------------------------------------------------- 排版工具

def make_helpers(Document, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, OxmlElement, qn, Mm, Pt):
    def set_run_font(run, name, size, bold=False):
        """中文用 name，英文和阿拉伯数字一律用 Times New Roman。

        w:eastAsia 管中文字符，w:ascii / w:hAnsi 管西文，Word 按字符所属
        Unicode 区段自动分流，所以一个 run 就够，不必按字符拆段。只设
        run.font.name 会把 ascii / hAnsi 也写成中文字体名——那样英文和数字
        就被迫用中文宋体渲染了，所以这里要显式盖回来。
        """
        run.font.name = name
        run.font.size = size
        run.bold = bold
        rPr = run._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:eastAsia"), name)
        rFonts.set(qn("w:ascii"), FONT_WESTERN)
        rFonts.set(qn("w:hAnsi"), FONT_WESTERN)
        # 复杂文种（阿拉伯文之类）也用西文字体，免得 Word 回退到默认值
        rFonts.set(qn("w:cs"), FONT_WESTERN)

    def _set_ind_attr(par, attr, value):
        pPr = par._p.get_or_add_pPr()
        ind = pPr.find(qn("w:ind"))
        if ind is None:
            ind = OxmlElement("w:ind")
            pPr.append(ind)
        ind.set(qn(attr), str(value))

    def add_par(doc, text="", *, font=FONT_BODY, size=None, align=None,
                indent_chars=INDENT_CHARS, left_chars=0.0, right_chars=0.0,
                hanging_chars=0.0, bold=False):
        if size is None:
            size = Pt(SIZE_BODY)
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(LINE_PITCH_PT)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        if align is not None:
            p.alignment = align
        # 先用 pt 值把 w:ind 建出来（保证 pPr 里顺序合法），再补字符值
        # hanging_chars：续行比首行多缩进的字数（附件说明用）
        pf.first_line_indent = Pt((indent_chars - hanging_chars) * SIZE_BODY)
        if left_chars or hanging_chars:
            pf.left_indent = Pt((left_chars + hanging_chars) * SIZE_BODY)
        if right_chars:
            pf.right_indent = Pt(right_chars * SIZE_BODY)
        if indent_chars and not hanging_chars:
            _set_ind_attr(p, "w:firstLineChars", int(indent_chars * 100))
        if hanging_chars:
            _set_ind_attr(p, "w:hangingChars", int(hanging_chars * 100))
        if left_chars:
            _set_ind_attr(p, "w:leftChars", int(left_chars * 100))
        if right_chars:
            _set_ind_attr(p, "w:rightChars", int(right_chars * 100))
        run = p.add_run(text)
        set_run_font(run, font, size, bold=bold)
        return p

    def blank_line(doc, count=1):
        for _ in range(count):
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(LINE_PITCH_PT)
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)

    def setup_section(doc):
        s = doc.sections[0]
        s.page_width = Mm(210)
        s.page_height = Mm(297)
        s.top_margin = Mm(MM_TOP)
        s.bottom_margin = Mm(MM_BOTTOM)
        s.left_margin = Mm(MM_LEFT)
        s.right_margin = Mm(MM_RIGHT)
        s.footer_distance = Mm(28)
        sectPr = s._sectPr
        docGrid = sectPr.find(qn("w:docGrid"))
        if docGrid is None:
            docGrid = OxmlElement("w:docGrid")
            sectPr.append(docGrid)
        docGrid.set(qn("w:type"), "lines")
        docGrid.set(qn("w:linePitch"), str(LINE_PITCH_TWIPS))
        return s

    def _append_field(par, instr, placeholder="1"):
        def field_run(build):
            r = par.add_run()
            build(r._r)
            return r

        def begin(el):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "begin")
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = " %s " % instr
            el.append(fc)
            el.append(it)

        def separate(el):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "separate")
            el.append(fc)
            t = OxmlElement("w:t")
            t.text = placeholder
            el.append(t)

        def end(el):
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "end")
            el.append(fc)

        return [field_run(begin), field_run(separate), field_run(end)]

    def add_page_numbers(doc, section, hide_first_page=False):
        """写入奇偶页页码；返回失败项列表，空列表表示全部成功。

        「单/双页居右、居左空一字」里的「一字」按正文字号（3 号 = 16 pt）计宽，
        与首行缩进同一把尺子；页码本身是 4 号字，两者口径不同，此处取正文口径。
        """
        problems = []
        try:
            doc.settings.odd_and_even_pages_header_footer = True
        except Exception as exc:
            problems.append("奇偶页页脚开关设置失败：%s" % exc)

        def fill(par, align, edge_pt):
            par.alignment = align
            pf = par.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            if align == WD_ALIGN_PARAGRAPH.RIGHT:
                pf.right_indent = Pt(edge_pt)
            else:
                pf.left_indent = Pt(edge_pt)
            runs = [par.add_run("— ")]
            runs.extend(_append_field(par, "PAGE"))
            runs.append(par.add_run(" —"))
            for r in runs:
                set_run_font(r, FONT_PAGE, Pt(SIZE_PAGE))

        try:
            odd = section.footer.paragraphs[0]
            fill(odd, WD_ALIGN_PARAGRAPH.RIGHT, SIZE_BODY)  # 单页：居右空一字
        except Exception as exc:
            problems.append("单页页码写入失败：%s" % exc)

        try:
            section.even_page_footer.is_linked_to_previous = False
            even = section.even_page_footer.paragraphs[0]
            fill(even, WD_ALIGN_PARAGRAPH.LEFT, SIZE_BODY)  # 双页：居左空一字
        except Exception as exc:
            problems.append("双页页码写入失败：%s" % exc)

        if hide_first_page:
            try:
                section.different_first_page_header_footer = True
            except Exception as exc:
                problems.append("首页不编页码设置失败：%s" % exc)
        return problems

    return set_run_font, add_par, blank_line, setup_section, add_page_numbers


def text_width_chars(s: str) -> float:
    """估算字符宽度：全角算 1 字，半角算 0.5 字。"""
    # 全角判定：CJK 区与全角标点，外加 × 号这类在公文里按全角计宽的符号
    wide_symbols = {0xD7, 0x2014, 0x2018, 0x2019, 0x201C, 0x201D, 0x2026, 0x00B7}
    w = 0.0
    for ch in s:
        if ord(ch) > 0x2E7F or ord(ch) in wide_symbols:
            w += 1.0
        else:
            w += 0.5
    return w


def _needs_space(prev: str, nxt: str) -> bool:
    """段内换行合并时：英文或数字相接补一个空格，中文直接相连。"""
    if not prev or not nxt:
        return False
    return prev.isascii() and nxt.isascii() and prev.isalnum() and nxt.isalnum()


# 标题标记 → 该标记占用的字符数（含后面的空格）
_HEAD_PREFIX_LEN = {"h1": 2, "h2": 3, "h3": 4, "h4": 5}

# 各级标题在文本里该配的编号，用来做「标记 ↔ 编号」配套检查
_CN_ORD_RE = re.compile(r"^[一二三四五六七八九十]{1,3}、")
_CN_PAREN_RE = re.compile(r"^（[一二三四五六七八九十]{1,3}）")
_AR_ORD_RE = re.compile(r"^\d{1,2}[.、]")
_AR_PAREN_RE = re.compile(r"^（\d{1,2}）")

_HEAD_NUMBERING = {
    "h1": ("#", "一、", _CN_ORD_RE),
    "h2": ("##", "（一）", _CN_PAREN_RE),
    "h3": ("###", "1.", _AR_ORD_RE),
    "h4": ("####", "（1）", _AR_PAREN_RE),
}

# 列表项前缀：短横线类、阿拉伯数字类、全角括号类。
# 数字后面必须跟空白、中文或行尾，否则「1.5 倍」这种小数会被误判成列表项。
_LIST_MARK_RE = re.compile(
    r"^(?:[-*+]\s+"
    r"|\d{1,2}[.、](?=[\s\u4e00-\u9fff]|$)"
    r"|（\d{1,2}）"
    r"|（[一二三四五六七八九十]{1,3}）)"
)


def parse_body(text: str):
    """空行分段，另加两条显式换行规则：

    - 行尾一个反斜杠 → 本行单独成段（强制换行，不跟下一行合并）；
    - `- `/`* `/`+ `、`1. `、`（1）`、`（一）` 开头的行 → 列表项，各占一段。

    `#`~`####` 开头是标题段。`\\#`、`\\-` 这类转义行按普通正文处理，
    不会变成标题或列表项。整份文件一个空行都没有时，退回「一行一段」
    （否则会把整篇正文合并成一大段）。
    """
    blocks: list[tuple[str, str]] = []
    buf: list[str] = []
    lines = text.splitlines()
    has_blank = any(not ln.strip() for ln in lines)

    def flush():
        if buf:
            blocks.append(("p", "".join(buf)))
            buf.clear()

    for raw in lines:
        line = raw.rstrip()

        # 行尾反斜杠 = 强制换行。先剥掉，免得它混进正文。
        # 双反斜杠是转义的「字面反斜杠」，不当换行符。
        hard_break = line.endswith("\\") and not line.endswith("\\\\")
        if hard_break:
            line = line[:-1].rstrip()

        if not line.strip():
            flush()
            continue

        # 转义行：吃掉反斜杠后按普通正文走，不再参与标题和列表识别
        escaped = len(line) >= 2 and line[0] == "\\" and line[1] in "#-*+"
        if escaped:
            line = line[1:]

        kind = ""
        if not escaped:
            for candidate in ("h4", "h3", "h2", "h1"):
                prefix = line[: _HEAD_PREFIX_LEN[candidate]]
                if prefix == "#" * int(candidate[1]) + " ":
                    kind = candidate
                    break
        if kind:
            flush()
            blocks.append((kind, line[_HEAD_PREFIX_LEN[kind]:].strip()))
            continue

        piece = line.strip()

        if not escaped and _LIST_MARK_RE.match(piece):
            flush()
            blocks.append(("p", _LIST_MARK_RE.sub("", piece, count=1).strip()))
            continue

        if hard_break:
            flush()
            blocks.append(("p", piece))
            continue

        if has_blank:
            if buf and _needs_space(buf[-1][-1:], piece[:1]):
                buf.append(" ")
            buf.append(piece)
        else:
            buf.append(piece)
            flush()
    flush()
    return blocks


def heading_warnings(blocks) -> list[str]:
    """层级标记和文本里的编号是否配套。只提醒，不阻断生成。

    `## 一、大作文` 这种写错层级的，会在这里被点出来——写错不会报错，
    但字体全错，是最难靠肉眼看出来的问题。
    """
    found: list[str] = []
    for kind, text in blocks:
        if kind not in _HEAD_NUMBERING:
            continue
        mark, want_label, own = _HEAD_NUMBERING[kind]
        head = text.strip()
        if own.match(head):
            continue
        for other_kind, (_mark, other_label, other_re) in _HEAD_NUMBERING.items():
            if other_kind != kind and other_re.match(head):
                found.append(
                    "「%s」写在 %s 标记下，用的却是「%s」的编号；%s 应该配「%s」"
                    % (head[:24], mark, other_label, mark, want_label))
                break
    return found


# ---------------------------------------------------------------- 输入校验

def _valid_date(s: str) -> bool:
    m = re.match(r"^(\d{4})\s*年\s*(\d{1,2})\s*月\s*(\d{1,2})\s*日$", s.strip())
    if not m:
        return False
    y, mo, d = (int(x) for x in m.groups())
    try:
        datetime.date(y, mo, d)
    except ValueError:
        return False
    return True


def strip_attach_prefix(s: str) -> str:
    s = s.strip()
    for prefix in ("附件：", "附件:"):
        if s.startswith(prefix):
            return s[len(prefix):].strip()
    return s


def validate_args(args) -> tuple[list[str], list[str]]:
    """返回 (阻断性错误, 提醒)。错误退出码 1，提醒只打印。"""
    errors: list[str] = []
    notices: list[str] = []

    if not args.out.lower().endswith(".docx"):
        errors.append("--out 必须以 .docx 结尾：%s" % _display_path(args.out))
    if not (args.title or "").strip():
        errors.append("--title 不能是空白")
    if args.date and not _valid_date(args.date):
        errors.append("--date 应为「YYYY 年 M 月 D 日」且日期真实存在：%s" % args.date)
    if args.to is not None and not args.to.strip():
        errors.append("--to 不能是空白（不要主送机关就别传这个参数）")
    elif args.to and not args.to.rstrip().endswith("："):
        errors.append("--to 应以全角冒号「：」结尾（主送机关后是冒号）：%s" % args.to)
    if args.body and not os.path.isfile(args.body):
        errors.append("正文文件不存在：%s" % _display_path(args.body))
    for item in args.attach or []:
        if item.strip().startswith(("附件：", "附件:")):
            notices.append("--attach 自带的「附件：」已去掉，避免重复：%s" % item)
    return errors, notices


# ---------------------------------------------------------------- 主流程

def _display_path(path: str) -> str:
    """日志里用的短路径：能相对当前目录就相对，否则只留文件名。"""
    try:
        rel = os.path.relpath(os.path.abspath(path))
    except (ValueError, OSError):
        return os.path.basename(path)
    if rel.startswith(".."):
        return os.path.basename(path)
    return rel


def build_document(args):
    (Document, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING,
     OxmlElement, qn, Mm, Pt) = _import_docx()
    (set_run_font, add_par, blank_line,
     setup_section, add_page_numbers) = make_helpers(
        Document, WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, OxmlElement, qn, Mm, Pt)

    doc = Document()
    section = setup_section(doc)

    # 标题：上方留 2 行空白，2 号方正小标宋简体，居中
    blank_line(doc, 2)
    add_par(doc, args.title, font=FONT_TITLE, size=Pt(SIZE_TITLE),
            align=WD_ALIGN_PARAGRAPH.CENTER, indent_chars=0)

    # 主送机关：标题下空 1 行，顶格
    if args.to:
        blank_line(doc, 1)
        add_par(doc, args.to, font=FONT_BODY, indent_chars=0)

    # 正文：主送机关下空 1 行
    body_text = ""
    if args.body:
        try:
            with open(args.body, "r", encoding="utf-8-sig") as fh:
                body_text = fh.read()
        except UnicodeDecodeError:
            print("正文文件不是 UTF-8 编码：%s（用记事本「另存为」选 UTF-8 再试）"
                  % _display_path(args.body), file=sys.stderr)
            raise SystemExit(1)
        except OSError as exc:
            print("读不到正文文件 %s：%s" % (_display_path(args.body), _os_error_text(exc)),
                  file=sys.stderr)
            raise SystemExit(1)
        if not body_text.strip():
            print("正文文件是空的：%s" % _display_path(args.body), file=sys.stderr)
            raise SystemExit(1)
    blocks = parse_body(body_text)
    if blocks:
        blank_line(doc, 1)
    for kind, text in blocks:
        if kind == "h1":
            add_par(doc, text, font=FONT_H1)
        elif kind == "h2":
            add_par(doc, text, font=FONT_H2)
        elif kind in ("h3", "h4"):
            # 三、四级标题字体同正文，区别只在于独立成段
            add_par(doc, text, font=FONT_H3)
        else:
            add_par(doc, text, font=FONT_BODY)

    # 附件说明：正文下空 1 行，左空 2 字；续行与附件名称首字对齐
    # 先滤掉空白项再决定要不要空行：整组都是空白时不该在正文后留一个空段
    attach_items = [text for text in
                    (strip_attach_prefix(item) for item in (args.attach or []))
                    if text]
    if attach_items:
        blank_line(doc, 1)
        for idx, text in enumerate(attach_items):
            if idx == 0:
                # 「附件：」占 3 字，悬挂 3 字让第 2 行起对齐到「1.」
                add_par(doc, "附件：%s" % text, font=FONT_BODY,
                        indent_chars=0, left_chars=INDENT_CHARS, hanging_chars=3.0)
            else:
                # 保持与第 1 条的「1.」对齐；悬挂 1.5 字让折行退到名称首字附近
                add_par(doc, text, font=FONT_BODY,
                        indent_chars=0, left_chars=INDENT_CHARS + 3.0,
                        hanging_chars=1.5)

    # 发文机关署名 + 成文日期：附件说明下空 1 行；日期右空 4 字，署名居中于日期
    if args.sign or args.date:
        blank_line(doc, 1)
        date_w = text_width_chars(args.date) if args.date else 0.0
        if args.sign:
            if args.date:
                # 以成文日期为准居中：署名比日期宽时向右伸出，窄时向内收
                right = 4.0 + (date_w - text_width_chars(args.sign)) / 2.0
                right = max(0.0, right)
            else:
                right = 4.0
            add_par(doc, args.sign, font=FONT_BODY, align=WD_ALIGN_PARAGRAPH.RIGHT,
                    indent_chars=0, right_chars=right)
        if args.date:
            add_par(doc, args.date, font=FONT_BODY, align=WD_ALIGN_PARAGRAPH.RIGHT,
                    indent_chars=0, right_chars=4.0)

    # 附注：成文日期下 1 行，左空 2 字，圆括号
    if args.note:
        note = args.note
        if not (note.startswith("（") and note.endswith("）")):
            note = "（%s）" % note.strip("（）()")
        add_par(doc, note, font=FONT_BODY, indent_chars=0, left_chars=INDENT_CHARS)

    page_problems = add_page_numbers(doc, section, hide_first_page=args.hide_first_page_number)
    return doc, blocks, page_problems


class _ArgumentParser(argparse.ArgumentParser):
    """参数错误统一退出码 1，避免与「缺字体 = 2」混淆。"""

    def error(self, message):
        self.print_usage(sys.stderr)
        self.exit(1, "%s: 参数错误：%s\n" % (self.prog, message))


def verify_output(path: str) -> list[str]:
    """生成后自检：把交付清单里能自动查的项目过一遍。

    查的是「写进去的东西对不对」——页边距、固定行距、英文数字字体的分流、
    一级标题的字体。查不了的项目直接跳过，不假装查过。
    """
    Document, _wa, _wl, _ox, qn, _mm, _pt = _import_docx()
    try:
        doc = Document(path)
    except Exception as exc:
        return ["打不开生成的文档：%s" % _error_text(exc)]

    problems: list[str] = []
    section = doc.sections[0]

    for name, got, want in (("上", section.top_margin.mm, MM_TOP),
                            ("下", section.bottom_margin.mm, MM_BOTTOM),
                            ("左", section.left_margin.mm, MM_LEFT),
                            ("右", section.right_margin.mm, MM_RIGHT)):
        if abs(got - want) > 0.2:
            problems.append("页边距%s %.1f mm，应为 %d mm" % (name, got, want))

    def _runs(paragraphs):
        for par in paragraphs:
            for run in par.runs:
                if not run.text.strip():
                    continue
                props = run._element.rPr
                fonts = props.rFonts if props is not None else None
                if fonts is None:
                    yield par, run.text, None, None, None
                else:
                    yield (par, run.text,
                           fonts.get(qn("w:ascii")),
                           fonts.get(qn("w:hAnsi")),
                           fonts.get(qn("w:eastAsia")))

    bad_west: list[str] = []
    for _par, text, ascii_font, h_ansi, _east in _runs(doc.paragraphs):
        if ascii_font != FONT_WESTERN or h_ansi != FONT_WESTERN:
            bad_west.append("「%s」用了 %s" % (text[:12], ascii_font))
    for footer in (section.footer, section.even_page_footer):
        for _par, text, ascii_font, h_ansi, _east in _runs(footer.paragraphs):
            if ascii_font != FONT_WESTERN or h_ansi != FONT_WESTERN:
                bad_west.append("页脚「%s」用了 %s" % (text[:12], ascii_font))
    if bad_west:
        problems.append("%d 处英文数字不是 %s（例如%s）"
                        % (len(bad_west), FONT_WESTERN, "、".join(bad_west[:3])))

    for _par, text, _a, _h, east in _runs(doc.paragraphs):
        if _CN_ORD_RE.match(text.strip()) and east != FONT_H1:
            problems.append("一级标题「%s」用了 %s，应为 %s"
                            % (text.strip()[:16], east, FONT_H1))

    for par in doc.paragraphs:
        if not par.text.strip():
            continue
        spacing = par.paragraph_format.line_spacing
        if spacing is None:
            continue
        pitch = getattr(spacing, "pt", None)
        if pitch is None:
            problems.append("「%s」不是固定行距" % par.text.strip()[:14])
            break
        if abs(pitch - LINE_PITCH_PT) > 0.5:
            problems.append("「%s」行距 %.1f 磅，应为 %d 磅"
                            % (par.text.strip()[:14], pitch, LINE_PITCH_PT))
            break

    return problems


def main(argv=None):
    _setup_console()
    parser = _ArgumentParser(
        description="按 GB/T 9704-2012 版式生成公文 Word（无红头）")
    parser.add_argument("--check-fonts", action="store_true", help="只检查字体")
    parser.add_argument("--force", "--allow-missing-fonts", dest="force",
                        action="store_true",
                        help="缺字体也继续生成（只放行，不替换字体名）")
    parser.add_argument("--out", help="输出 .docx 路径")
    parser.add_argument("--title", help="标题")
    parser.add_argument("--to", help="主送机关，如「各处室：」")
    parser.add_argument("--body",
                        help="正文纯文本文件（# / ## / ### / #### 层级；"
                             "- / 1. / （1） / （一） 开头的行各成一段；"
                             "行尾 \\ 强制换行；\\# 转义行首井号）")
    parser.add_argument("--sign", help="发文机关署名")
    parser.add_argument("--date", help="成文日期，如 2025年10月25日")
    parser.add_argument("--attach", action="append", help="附件说明，可重复")
    parser.add_argument("--note", help="附注（圆括号内容）")
    parser.add_argument("--hide-first-page-number", action="store_true",
                        help="首页不编页码（发文惯例）")
    parser.add_argument("--verify", action="store_true",
                        help="生成后自检字体分流与版式，有问题退出码 6")
    args = parser.parse_args(argv)

    if args.check_fonts:
        missing, detect_error = find_missing_fonts()
        if detect_error:
            print("字体检测不可用：%s" % detect_error, file=sys.stderr)
            print("无法确认字体是否齐全。确认已装好后，加 --force（别名 "
                  "--allow-missing-fonts）可跳过检查。", file=sys.stderr)
            return 5
        if not missing:
            print("字体齐全：%s" % " / ".join(font for font, _ in REQUIRED_FONTS))
            return 0
        print("缺少以下公文字体（共 %d 项）：" % len(missing))
        for font, role in missing:
            print("  - %-18s（用于%s）" % (font, role))
        print("请先安装上述字体；临时可用 --force（别名 --allow-missing-fonts）继续，"
              "但脚本不会替换字体名，显示效果由打开文档的机器决定。")
        return 2

    if not args.out:
        parser.error("生成文档需要 --out")

    errors, notices = validate_args(args)
    for note in notices:
        print("提醒：%s" % note)
    if errors:
        for err in errors:
            print("参数错误：%s" % err, file=sys.stderr)
        return 1

    missing, detect_error = find_missing_fonts()
    if detect_error:
        print("字体检测不可用：%s" % detect_error, file=sys.stderr)
        if not args.force:
            print("无法确认字体是否齐全，已停止。确认已装好后加 --force 继续。",
                  file=sys.stderr)
            return 5
        print("已按 --force 继续：字体检查结果不可信，请自行确认字体已装。",
              file=sys.stderr)
    elif missing:
        print("缺少以下公文字体（共 %d 项）：" % len(missing), file=sys.stderr)
        for font, role in missing:
            print("  - %-18s（用于%s）" % (font, role), file=sys.stderr)
        if not args.force:
            print("未获允许，已停止。装上字体后重跑，或明确同意后加 --force。", file=sys.stderr)
            return 2
        print("已按 --force 继续生成：文档内仍写规范字体名，缺失字体由 Word/WPS "
              "打开时替代显示，实际效果不可控。", file=sys.stderr)

    doc, blocks, page_problems = build_document(args)

    # 层级配错不会报错，但字体全错——落盘前先提醒一次
    for warning in heading_warnings(blocks):
        print("提醒：%s" % warning, file=sys.stderr)

    out_dir = os.path.dirname(os.path.abspath(args.out))
    tmp_path = ""

    def _cleanup_tmp() -> None:
        if not tmp_path:
            return
        try:
            os.unlink(tmp_path)
        except Exception as cleanup_exc:
            # 清理本身出错只提醒，绝不替换掉真正的异常（下面 except 块里会原样抛出）
            print("提醒：临时文件没清掉：%s（%s）"
                  % (_display_path(tmp_path), _error_text(cleanup_exc)),
                  file=sys.stderr)

    try:
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)
        # 唯一临时名：同一目录里并发生成两份文档时不会互相覆盖
        fd, tmp_path = tempfile.mkstemp(
            prefix=os.path.basename(args.out) + ".", suffix=".tmp",
            dir=out_dir or None)
        os.close(fd)
        doc.save(tmp_path)
        os.replace(tmp_path, args.out)
    except OSError as exc:
        print("写入失败：%s（目标 %s）" % (_os_error_text(exc), _display_path(args.out)),
              file=sys.stderr)
        _cleanup_tmp()
        return 1
    except BaseException:
        # 不是 OSError 的意外（python-docx 内部异常等）也要清掉临时文件，
        # 再把异常原样抛出去——吞掉的话真正的 bug 就再也没人知道了。
        _cleanup_tmp()
        raise

    print("已生成：%s" % _display_path(args.out))
    print("版式：A4，页边距 上%d/下%d/左%d/右%d mm，固定行距 %d 磅，每页 22 行"
          % (MM_TOP, MM_BOTTOM, MM_LEFT, MM_RIGHT, LINE_PITCH_PT))
    print("正文段数：%d" % len(blocks))
    print("字体：标题 %s %dpt｜正文 %s %dpt｜页码 %s %dpt｜英文数字 %s"
          % (FONT_TITLE, SIZE_TITLE, FONT_BODY, SIZE_BODY, FONT_PAGE, SIZE_PAGE,
             FONT_WESTERN))
    if detect_error:
        print("注意：字体未做检查（检测不可用），文档内仍按规范字体名写入。")
    elif missing:
        print("注意：下列字体缺失，文档内仍按规范字体名写入（打开时由系统替代）：%s"
              % "、".join(f for f, _ in missing))
    if page_problems:
        print("警告：页码设置未全部成功，请检查页脚：", file=sys.stderr)
        for problem in page_problems:
            print("  - %s" % problem, file=sys.stderr)
        return 4

    if args.verify:
        problems = verify_output(args.out)
        if problems:
            print("自检发现 %d 个问题：" % len(problems), file=sys.stderr)
            for problem in problems:
                print("  - %s" % problem, file=sys.stderr)
            return 6
        print("自检通过：页边距、固定行距、中英文字体分流、一级标题字体均符合规范。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
